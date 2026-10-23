"""
Proje maliyet kalemleri ve bütçe uçları (ADIM 7.5.3). Proje yöneticisi (veya project:manage) yönetir.
Kalemler proje kapanana kadar değiştirilebilir. Belge (fatura vb.) öğrenme hattına girmez; nesne deposunda ayrı saklanır.
"""
import uuid
from datetime import date
from decimal import Decimal
from typing import Optional
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import Response
from pydantic import BaseModel, Field
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.security import CurrentUser, get_current_user
from app.db.models import ProjectCostItem, Task
from app.db.session import get_db
from app.services import costs, storage
from app.services import projects as psvc

router = APIRouter(prefix="/projects", tags=["Proje Maliyeti"])

CLOSED = ("COMPLETED", "CANCELLED")
MAX_ATTACHMENT = 20 * 1024 * 1024
ATTACHMENT_TYPES = {"pdf": "application/pdf", "png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg",
                    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"}


class ItemIn(BaseModel):
    category: str = Field(pattern="^(MATERYAL|DANISMANLIK|EGITIM|LISANS|SEYAHAT|DONANIM|DIGER)$")
    description: str = Field(min_length=2, max_length=500)
    quantity: Decimal = Field(gt=0, le=1_000_000)
    unit: Optional[str] = Field(None, max_length=24)
    unit_price_try: Decimal = Field(ge=0, le=1_000_000_000)
    incurred_on: Optional[date] = None
    task_id: Optional[uuid.UUID] = None


async def _project(db, user, project_id, editable=True):
    project = await psvc.load(db, user, project_id, manage=True)
    if editable and project.status in CLOSED:
        raise HTTPException(status_code=409, detail="Kapanmış projenin maliyet kalemleri değiştirilemez.")
    return project


async def _check_task(db, project, task_id):
    if task_id:
        task = await db.get(Task, task_id)
        if not task or task.project_id != project.id:
            raise HTTPException(status_code=422, detail="İş adımı bu projeye ait değil.")


async def _item(db, project, item_id) -> ProjectCostItem:
    item = await db.get(ProjectCostItem, item_id)
    if not item or item.project_id != project.id:
        raise HTTPException(status_code=404, detail="Maliyet kalemi bulunamadı.")
    return item


def _amount(q: Decimal, p: Decimal) -> Decimal:
    return (q * p).quantize(Decimal("0.01"))


@router.get("/{project_id}/cost-items")
async def list_items(project_id: uuid.UUID, user: CurrentUser = Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    project = await _project(db, user, project_id, editable=False)
    return (await costs.project_costs(db, project))["items"]


@router.post("/{project_id}/cost-items", status_code=201)
async def add_item(project_id: uuid.UUID, payload: ItemIn, user: CurrentUser = Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    project = await _project(db, user, project_id)
    await _check_task(db, project, payload.task_id)
    item = ProjectCostItem(tenant_id=project.tenant_id, project_id=project.id, task_id=payload.task_id, category=payload.category,
                           description=payload.description.strip(), quantity=payload.quantity, unit=(payload.unit or None),
                           unit_price_try=payload.unit_price_try, amount_try=_amount(payload.quantity, payload.unit_price_try),
                           incurred_on=payload.incurred_on or date.today(), created_by_actor_id=user.actor_id)
    db.add(item)
    await db.commit()
    return costs.item_out(item)


@router.put("/{project_id}/cost-items/{item_id}")
async def update_item(project_id: uuid.UUID, item_id: uuid.UUID, payload: ItemIn, user: CurrentUser = Depends(get_current_user),
                      db: AsyncSession = Depends(get_db)):
    project = await _project(db, user, project_id)
    item = await _item(db, project, item_id)
    await _check_task(db, project, payload.task_id)
    item.category, item.description, item.quantity = payload.category, payload.description.strip(), payload.quantity
    item.unit, item.unit_price_try, item.task_id = (payload.unit or None), payload.unit_price_try, payload.task_id
    item.amount_try = _amount(payload.quantity, payload.unit_price_try)
    item.incurred_on = payload.incurred_on or item.incurred_on
    await db.commit()
    return costs.item_out(item)


@router.delete("/{project_id}/cost-items/{item_id}")
async def delete_item(project_id: uuid.UUID, item_id: uuid.UUID, user: CurrentUser = Depends(get_current_user),
                      db: AsyncSession = Depends(get_db)):
    project = await _project(db, user, project_id)
    item = await _item(db, project, item_id)
    if item.attachment_path:
        try:
            await storage.delete_object(item.attachment_path)
        except Exception:
            pass
    await db.delete(item)
    await db.commit()
    return {"success": True}


@router.post("/{project_id}/cost-items/{item_id}/attachment")
async def upload_attachment(project_id: uuid.UUID, item_id: uuid.UUID, file: UploadFile = File(...),
                            user: CurrentUser = Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    project = await _project(db, user, project_id)
    item = await _item(db, project, item_id)
    ext = (file.filename or "").rsplit(".", 1)[-1].lower()
    if ext not in ATTACHMENT_TYPES:
        raise HTTPException(status_code=422, detail="Belge PDF, PNG, JPG, XLSX veya DOCX olmalı.")
    data = await file.read()
    if len(data) > MAX_ATTACHMENT:
        raise HTTPException(status_code=413, detail="Belge en fazla 20 MB olabilir.")
    key = f"cost-items/{project.id}/{item.id}.{ext}"
    if item.attachment_path and item.attachment_path != key:
        try:
            await storage.delete_object(item.attachment_path)
        except Exception:
            pass
    await storage.put_object(key, data, ATTACHMENT_TYPES[ext])
    item.attachment_path, item.attachment_name = key, (file.filename or f"belge.{ext}")[:255]
    await db.commit()
    return costs.item_out(item)


@router.get("/{project_id}/cost-items/{item_id}/attachment")
async def download_attachment(project_id: uuid.UUID, item_id: uuid.UUID, user: CurrentUser = Depends(get_current_user),
                              db: AsyncSession = Depends(get_db)):
    project = await _project(db, user, project_id, editable=False)
    item = await _item(db, project, item_id)
    if not item.attachment_path:
        raise HTTPException(status_code=404, detail="Bu kalemin belgesi yok.")
    data = await storage.get_object(item.attachment_path)
    ext = item.attachment_path.rsplit(".", 1)[-1]
    return Response(content=data, media_type=ATTACHMENT_TYPES.get(ext, "application/octet-stream"),
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(item.attachment_name or 'belge')}"})


class BudgetIn(BaseModel):
    budget_try: Optional[Decimal] = Field(None, ge=0, le=100_000_000_000)
    budget_note: Optional[str] = Field(None, max_length=2000)


@router.put("/{project_id}/budget")
async def set_budget(project_id: uuid.UUID, payload: BudgetIn, user: CurrentUser = Depends(get_current_user),
                     db: AsyncSession = Depends(get_db)):
    project = await _project(db, user, project_id)
    project.budget_try, project.budget_note = payload.budget_try, (payload.budget_note or None)
    await db.commit()
    return {"budget_try": float(project.budget_try) if project.budget_try is not None else None, "budget_note": project.budget_note}


# ---------------------------------------------------------------- kapanış raporu (ADIM 7.5.4)
@router.get("/{project_id}/costs/export.xlsx")
async def export_costs(project_id: uuid.UUID, user: CurrentUser = Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    """Proje maliyet raporu (Excel): Özet, İş adımları, Maliyet kalemleri, LLM kırılımı. TL, KDV hariç."""
    import io
    from datetime import datetime, timezone

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    project = await _project(db, user, project_id, editable=False)
    c = await costs.project_costs(db, project)
    titles = {t["task_id"]: t["title"] for t in c["tasks"]}
    money, head = '#,##0.00 "₺"', PatternFill("solid", fgColor="E9E3FB")
    basis = {"ACTUAL": "Gerçekleşen", "RUNNING": "Sürüyor", "PLANNED": "Plan (tahmin)"}

    wb = Workbook()

    def sheet(ws, header, rows, money_cols=(), widths=None):
        ws.append(header)
        for cell in ws[1]:
            cell.font, cell.fill, cell.alignment = Font(bold=True), head, Alignment(vertical="center", wrap_text=True)
        for r in rows:
            ws.append(r)
        for col in money_cols:
            for (cell,) in ws.iter_rows(min_row=2, min_col=col, max_col=col):
                cell.number_format = money
        for i, w in enumerate(widths or [], 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"

    t = c["totals"]
    ws = wb.active
    ws.title = "Özet"
    summary = [
        ("Proje", project.name), ("Durum", svc_status(project.status)),
        ("Tarih aralığı", f"{project.start_date or '-'} - {project.end_date or '-'}"),
        ("Bütçe", c["budget"]["amount"]), ("Bütçe notu", c["budget"]["note"] or ""),
        ("Gerçekleşen: personel", t["labor"]), ("Gerçekleşen: token (LLM)", t["token"]), ("Gerçekleşen: diğer kalemler", t["other"]),
        ("Gerçekleşen toplam", t["actual"]), ("Bütçe farkı (bütçe - gerçekleşen)", t["budget_variance"]),
        ("Planlanan maliyet (bütçe önerisi)", c["budget"]["suggested"]), ("Başlamamış adımların tahmini personel maliyeti", t["labor_planned"]),
        ("Fiyatı tanımsız LLM çağrısı", t["unpriced_calls"]), ("Harcanan RECOrion kredisi", t["credits"]),
        ("Para birimi", "TL, KDV hariç"), ("Rapor zamanı", datetime.now(timezone.utc).strftime("%d.%m.%Y %H:%M UTC")),
    ]
    sheet(ws, ["Kalem", "Değer"], summary, widths=[46, 40])
    for row in ws.iter_rows(min_row=2):
        if isinstance(row[1].value, float) and row[0].value not in ("Harcanan RECOrion kredisi",):
            row[1].number_format = money
    sheet(wb.create_sheet("İş adımları"),
          ["#", "İş adımı", "Personel", "Tür", "Durum", "Gün", "Dayanak", "Personel maliyeti", "Token (LLM)", "Çağrı", "Fiyatsız çağrı",
           "Kredi", "Diğer kalemler", "Toplam"],
          [[x["order_index"], x["title"], x["assignee"] or "-", "AI" if x["assignee_type"] == "AI_AGENT" else "İnsan" if x["assignee"] else "-",
            x["status"], x["labor"]["days"], basis.get(x["labor"]["basis"], x["labor"]["basis"]), x["labor"]["cost"], x["tokens"]["cost"],
            x["tokens"]["calls"], x["tokens"]["unpriced_calls"], x["tokens"]["credits"], x["other"],
            (0 if x["labor"]["basis"] == "PLANNED" else x["labor"]["cost"]) + x["tokens"]["cost"] + x["other"]] for x in c["tasks"]]
          + ([["", "Proje geneli (planlama, adıma bağlı olmayan)", "", "", "", "", "", 0, c["project_level_tokens"]["cost"],
               c["project_level_tokens"]["calls"], c["project_level_tokens"]["unpriced_calls"], c["project_level_tokens"]["credits"], 0,
               c["project_level_tokens"]["cost"]]] if c["project_level_tokens"] else []),
          money_cols=(8, 9, 13, 14), widths=[5, 40, 24, 8, 16, 7, 14, 16, 14, 8, 12, 10, 14, 14])
    sheet(wb.create_sheet("Maliyet kalemleri"),
          ["Tarih", "Kategori", "Açıklama", "İş adımı", "Miktar", "Birim", "Birim fiyat", "Tutar", "Belge"],
          [[str(i["incurred_on"]), i["category_label"], i["description"], titles.get(i["task_id"], "-") if i["task_id"] else "-",
            i["quantity"], i["unit"] or "", i["unit_price_try"], i["amount_try"], i["attachment_name"] or ""] for i in c["items"]],
          money_cols=(7, 8), widths=[12, 16, 44, 30, 9, 9, 14, 14, 26])
    sheet(wb.create_sheet("LLM kırılımı"),
          ["Model", "Çağrı", "Girdi token", "Çıktı token", "Token (LLM) TL", "Fiyatsız çağrı", "Kredi"],
          [[m["model_code"], m["calls"], m["tokens_in"], m["tokens_out"], m["cost"], m["unpriced_calls"], m["credits"]] for m in c["by_model"]],
          money_cols=(5,), widths=[26, 8, 14, 14, 16, 14, 10])
    buf = io.BytesIO()
    wb.save(buf)
    name = f"{project.slug or 'proje'}-maliyet-{datetime.now().strftime('%Y%m%d')}.xlsx"
    return Response(content=buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(name)}"})


def svc_status(status: str) -> str:
    return psvc.STATUS_LABELS.get(status, status)
