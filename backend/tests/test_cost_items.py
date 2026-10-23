"""
ADIM 7.5.3: Diğer maliyet kalemleri, belge ve proje bütçesi.

- Kalem: kategori, miktar × birim fiyat = tutar; isteğe bağlı iş adımı; yalnızca proje yöneticisi; kapanmış projede kilitli.
- Belge (fatura vb.) ayrı saklanır, öğrenme hattına girmez; türü ve boyutu denetlenir.
- Bütçe, bütçe önerisi (planlanan personel + token) ve bütçe farkı maliyet raporunda.
"""
import pytest

from app.services import storage
from tests.test_workflow import _exec, _ready_project, world  # noqa: F401  (fikstür)


@pytest.fixture
def store(monkeypatch):
    objects = {}

    async def put_object(key, data, content_type=None):
        objects[key] = data

    async def get_object(key):
        return objects[key]

    async def delete_object(key):
        objects.pop(key, None)

    monkeypatch.setattr(storage, "put_object", put_object)
    monkeypatch.setattr(storage, "get_object", get_object)
    monkeypatch.setattr(storage, "delete_object", delete_object)
    return objects


def test_cost_items_attachment_budget(client, db, world, store):
    w = world
    admin, member = w["admin"]["h"], w["member"]["h"]
    pid, human_id, dep_id = _ready_project(client, w, "Kalem Projesi")
    other_pid, other_human, _ = _ready_project(client, w, "Başka Proje")
    base = f"/api/v1/projects/{pid}/cost-items"

    # Ekleme, doğrulama ve yetki
    r = client.post(base, headers=admin, json={"category": "MATERYAL", "description": "Sunucu kabini", "quantity": 3, "unit": "adet", "unit_price_try": 150})
    assert r.status_code == 201, r.text
    mat = r.json()
    assert mat["amount_try"] == 450.0 and mat["category_label"] == "Materyal" and mat["incurred_on"]
    r = client.post(base, headers=admin, json={"category": "DANISMANLIK", "description": "Güvenlik danışmanlığı", "quantity": 1,
                                               "unit_price_try": 2000, "task_id": human_id})
    dan = r.json()
    assert r.status_code == 201 and dan["task_id"] == human_id
    assert client.post(base, headers=admin, json={"category": "YEMEK", "description": "x", "quantity": 1, "unit_price_try": 1}).status_code == 422
    assert client.post(base, headers=admin, json={"category": "DIGER", "description": "Başka projenin adımı", "quantity": 1,
                                                  "unit_price_try": 1, "task_id": other_human}).status_code == 422
    assert client.post(base, headers=member, json={"category": "DIGER", "description": "Üye", "quantity": 1, "unit_price_try": 1}).status_code == 403
    assert client.get(base, headers=member).status_code == 403

    # Güncelleme ve silme
    r = client.put(f"{base}/{mat['id']}", headers=admin, json={"category": "DONANIM", "description": "Sunucu kabini", "quantity": 2, "unit_price_try": 175})
    assert r.status_code == 200 and r.json()["amount_try"] == 350.0 and r.json()["category"] == "DONANIM"
    tmp = client.post(base, headers=admin, json={"category": "SEYAHAT", "description": "Silinecek", "quantity": 1, "unit_price_try": 99}).json()
    assert client.delete(f"{base}/{tmp['id']}", headers=admin).status_code == 200
    assert len(client.get(base, headers=admin).json()) == 2

    # Belge: tür denetimi, yükleme ve indirme (öğrenme hattına girmez: Proje Ekleri boş kalır)
    bad = client.post(f"{base}/{dan['id']}/attachment", headers=admin, files={"file": ("x.exe", b"MZ", "application/octet-stream")})
    assert bad.status_code == 422
    pdf = b"%PDF-1.4 fatura"
    up = client.post(f"{base}/{dan['id']}/attachment", headers=admin, files={"file": ("fatura-001.pdf", pdf, "application/pdf")})
    assert up.status_code == 200 and up.json()["has_attachment"] and up.json()["attachment_name"] == "fatura-001.pdf"
    down = client.get(f"{base}/{dan['id']}/attachment", headers=admin)
    assert down.status_code == 200 and down.content == pdf and "fatura-001.pdf" in down.headers["content-disposition"]
    assert client.get(f"/api/v1/projects/{pid}/documents", headers=admin).json() == []

    # Bütçe, öneri ve fark: uzman günlük 1000 ₺ (insan adımı 1 gün), AI modeli 1M girdi 10 ₺ / çıktı 40 ₺
    tid = w["tenant"]
    client.put(f"/api/v1/company/{tid}/users/{w['member']['id']}", headers=admin, json={"daily_cost_try": 1000})
    model_id = _exec(db, "SELECT preferred_model_id FROM actors WHERE id = %s", w["agent"])[0]
    _exec(db, "DELETE FROM llm_model_prices WHERE model_id = %s", model_id)
    _exec(db, "INSERT INTO llm_model_prices (model_id, price_in_per_1m_try, price_out_per_1m_try, valid_from) VALUES (%s, 10, 40, now() - interval '1 day')", model_id)
    try:
        assert client.put(f"/api/v1/projects/{pid}/budget", headers=member, json={"budget_try": 1}).status_code == 403
        assert client.put(f"/api/v1/projects/{pid}/budget", headers=admin, json={"budget_try": 5000, "budget_note": "Yönetim onaylı"}).status_code == 200
        c = client.get(f"/api/v1/projects/{pid}/costs", headers=admin).json()
        assert c["totals"]["other"] == 2350.0 and {x["category"] for x in c["by_category"]} == {"DONANIM", "DANISMANLIK"}
        human_row = next(t for t in c["tasks"] if t["task_id"] == human_id)
        assert human_row["other"] == 2000.0
        # Plan: insan adımı 1000 ₺ + 2 AI adımı × (40k×10 + 10k×40)/1M = 0,8 ₺ → 1001,60 ₺
        assert c["planned"]["labor"] == 1000.0 and c["planned"]["token"] == 1.6 and c["budget"]["suggested"] == 1001.6
        assert c["budget"]["amount"] == 5000.0 and c["budget"]["note"] == "Yönetim onaylı"
        assert c["totals"]["actual"] == 2350.0 and c["totals"]["budget_variance"] == 2650.0  # henüz başlamadı: yalnızca kalemler

        # Kapanış raporu (Excel, ADIM 7.5.4): 4 sayfa, özet tutarları raporla aynı; ekip üyesine kapalı
        import io

        from openpyxl import load_workbook

        assert client.get(f"/api/v1/projects/{pid}/costs/export.xlsx", headers=member).status_code == 403
        x = client.get(f"/api/v1/projects/{pid}/costs/export.xlsx", headers=admin)
        assert x.status_code == 200 and "maliyet" in x.headers["content-disposition"]
        wb = load_workbook(io.BytesIO(x.content))
        assert wb.sheetnames == ["Özet", "İş adımları", "Maliyet kalemleri", "LLM kırılımı"]
        summary = {r[0]: r[1] for r in wb["Özet"].iter_rows(min_row=2, values_only=True)}
        assert summary["Bütçe"] == 5000 and summary["Gerçekleşen: diğer kalemler"] == 2350 and summary["Bütçe farkı (bütçe - gerçekleşen)"] == 2650
        assert summary["Para birimi"] == "TL, KDV hariç"
        items = list(wb["Maliyet kalemleri"].iter_rows(min_row=2, values_only=True))
        assert len(items) == 2 and any(r[3] == "İnsan adımı" and r[7] == 2000 and r[8] == "fatura-001.pdf" for r in items)
        assert len(list(wb["İş adımları"].iter_rows(min_row=2, values_only=True))) == 3
    finally:
        _exec(db, "DELETE FROM llm_model_prices WHERE model_id = %s", model_id)
        _exec(db, "UPDATE actors SET daily_cost_try = NULL WHERE id = %s", w["member"]["id"])

    # Kapanmış projede kilit
    client.post(f"/api/v1/projects/{pid}/status", headers=admin, json={"action": "cancel"})
    assert client.post(base, headers=admin, json={"category": "DIGER", "description": "Geç", "quantity": 1, "unit_price_try": 1}).status_code == 409
    assert client.put(f"/api/v1/projects/{pid}/budget", headers=admin, json={"budget_try": 1}).status_code == 409
    assert client.get(base, headers=admin).status_code == 200  # okuma serbest
    client.post(f"/api/v1/projects/{other_pid}/status", headers=admin, json={"action": "cancel"})
