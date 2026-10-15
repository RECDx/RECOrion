import openpyxl

wb = openpyxl.load_workbook('documents/Sektör Meslek Nace Kodları_2026.03.xlsx', data_only=True)
ws = wb['Sheet0']

rows_with_any_data_col0_none = []
for r in range(2, ws.max_row + 1):
    c0 = ws.cell(r, 1).value
    all_vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
    if c0 is None and any(v is not None for v in all_vals):
        rows_with_any_data_col0_none.append((r, all_vals))

print(f"Rows where c0 is None but SOME column has data: {len(rows_with_any_data_col0_none)}")
for r, vals in rows_with_any_data_col0_none[:15]:
    print(f"Row {r}: {vals}")
