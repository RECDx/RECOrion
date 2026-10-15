import openpyxl

wb = openpyxl.load_workbook('documents/Sektör Meslek Nace Kodları_2026.03.xlsx', data_only=True)
ws = wb['Sheet0']

print("Row 2 sector name repr:", repr(ws.cell(2, 2).value))
print("Row 2 prof name repr:", repr(ws.cell(2, 4).value))
print("Row 1364 sector name repr (K):", repr(ws.cell(1364, 2).value))
print("Row 1364 prof name repr (K.01):", repr(ws.cell(1364, 4).value))
