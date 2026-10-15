import openpyxl

wb = openpyxl.load_workbook('documents/Sektör Meslek Nace Kodları_2026.03.xlsx', data_only=True)
print("All sheets:", wb.sheetnames)
for s in wb.worksheets:
    print(f"Sheet: {s.title}, state: {s.sheet_state}, max_row: {s.max_row}, max_col: {s.max_column}")
