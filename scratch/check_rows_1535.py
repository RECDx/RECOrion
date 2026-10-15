import openpyxl

wb = openpyxl.load_workbook('documents/Sektör Meslek Nace Kodları_2026.03.xlsx', data_only=True)
ws = wb['Sheet0']

print("Row 1:", [ws.cell(1, c).value for c in range(1, 9)])
for r in range(1520, 1540):
    print(f"Row {r}:", [ws.cell(r, c).value for c in range(1, 9)])
