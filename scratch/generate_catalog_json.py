import json
import os
import openpyxl

EXCEL_PATH = os.path.join("documents", "Sektör Meslek Nace Kodları_2026.03.xlsx")
OUTPUT_JSON = os.path.join("backend", "trade_ministry_catalog.json")

SECTOR_ICONS = {
    "A": "🌲",  # Ağaç İşleri
    "B": "🎭",  # Eğlence, Dinlenme, Organizasyon
    "C": "💻",  # Elektrik, Elektronik, Bilgisayar
    "D": "🌾",  # Gıda, Tarım
    "E": "👗",  # Giyim, Deri Ürün, Ev Tekstili, Dokuma
    "F": "🎁",  # Hediyelik Eşya, Eğitim, Basım, Fotoğraf, Çeşitli Mallar
    "G": "✂️",  # Kuaför, Berber, Temizlik, Spor, Kozmetik, Sağlık
    "H": "🚗",  # Metal, Otomotiv, Makine
    "I": "🛒",  # Pazar, Seyyar
    "J": "🚚",  # Ulaştırma Hizmetleri
    "K": "🏗️",  # Yapı Sanatları
}

def main():
    print(f"Reading: {EXCEL_PATH}")
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws = wb['Sheet0']
    
    sectors = {}
    professions = {}
    
    for r in range(2, ws.max_row + 1):
        c0 = ws.cell(r, 1).value
        c1 = ws.cell(r, 2).value
        c2 = ws.cell(r, 3).value
        c3 = ws.cell(r, 4).value
        
        if c0 is not None and c1 is not None:
            sc = str(c0).strip()
            sn = str(c1).strip()
            if sc and sn and sc not in sectors:
                sectors[sc] = {
                    "code": sc,
                    "name": sn,
                    "icon": SECTOR_ICONS.get(sc, "📁")
                }
                
        if c0 is not None and c2 is not None and c3 is not None:
            sc = str(c0).strip()
            pc = str(c2).strip()
            pn = str(c3).strip()
            if sc and pc and pn and pc not in professions:
                professions[pc] = {
                    "code": pc,
                    "sector_code": sc,
                    "name": pn
                }
                
    catalog_data = {
        "sectors": [sectors[k] for k in sorted(sectors.keys())],
        "professions": [professions[k] for k in sorted(professions.keys())]
    }
    
    print(f"Total Sectors: {len(catalog_data['sectors'])}")
    print(f"Total Professions: {len(catalog_data['professions'])}")
    
    with open(OUTPUT_JSON, "w", encoding="utf-8") as f:
        json.dump(catalog_data, f, ensure_ascii=False, indent=2)
        
    print(f"Catalog JSON successfully saved to {OUTPUT_JSON}!")

if __name__ == '__main__':
    main()
