import openpyxl

wb = openpyxl.load_workbook('documents/Sektör Meslek Nace Kodları_2026.03.xlsx', data_only=True)
ws = wb['Sheet0']

print("Dimensions:", ws.dimensions)
for col_idx in range(1, ws.max_column + 1):
    vals = [ws.cell(r, col_idx).value for r in range(1, 20)]
    print(f"Col {col_idx} (header: {vals[0]}): sample={vals[1:5]}")
