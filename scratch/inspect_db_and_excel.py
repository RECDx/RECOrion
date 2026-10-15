import asyncio
import asyncpg
import openpyxl

async def inspect():
    print("--- 1. DATABASE CHECK ---")
    conn = await asyncpg.connect('postgresql://recorion_admin:recorion_secret_pass_2026@localhost:5432/recorion_db')
    sectors = await conn.fetch('SELECT code, name FROM sectors ORDER BY code')
    print(f"Total sectors in DB: {len(sectors)}")
    for s in sectors:
        cnt = await conn.fetchval('SELECT count(*) FROM professions WHERE sector_code = $1', s['code'])
        print(f"  [{s['code']}] {s['name']}: {cnt} professions")
        
    profs_k = await conn.fetch('SELECT code, sector_code, name FROM professions WHERE sector_code = $1 ORDER BY code', 'K')
    print(f"Professions for K: {len(profs_k)}")
    for p in profs_k[:5]:
        print(f"  {p['code']} -> {p['name']}")
        
    await conn.close()

    print("\n--- 2. EXCEL CHECK ---")
    wb = openpyxl.load_workbook('documents/Sektör Meslek Nace Kodları_2026.03.xlsx', data_only=True)
    ws = wb['Sheet0']
    print(f"Sheet rows: {ws.max_row}, cols: {ws.max_column}")
    
    excel_sectors = {}
    excel_professions = {}
    
    for r in range(2, ws.max_row + 1):
        s_code = ws.cell(r, 1).value
        s_name = ws.cell(r, 2).value
        p_code = ws.cell(r, 3).value
        p_name = ws.cell(r, 4).value
        
        s_code_str = str(s_code).strip() if s_code is not None else ""
        s_name_str = str(s_name).strip() if s_name is not None else ""
        p_code_str = str(p_code).strip() if p_code is not None else ""
        p_name_str = str(p_name).strip() if p_name is not None else ""
        
        if s_code_str:
            if s_code_str not in excel_sectors:
                excel_sectors[s_code_str] = set()
            if s_name_str:
                excel_sectors[s_code_str].add(s_name_str)
                
        if p_code_str:
            if p_code_str not in excel_professions:
                excel_professions[p_code_str] = set()
            excel_professions[p_code_str].add((s_code_str, p_name_str))

    print(f"\nDistinct sector codes in Excel: {len(excel_sectors)}")
    for sc, names in sorted(excel_sectors.items()):
        print(f"  [{sc}]: {list(names)}")
        # Count distinct profession codes for this sector
        p_count = sum(1 for pc, items in excel_professions.items() if any(item[0] == sc for item in items))
        print(f"     -> has {p_count} distinct profession codes")

    print(f"\nTotal distinct profession codes in Excel: {len(excel_professions)}")

if __name__ == "__main__":
    asyncio.run(inspect())
