import openpyxl

wb = openpyxl.load_workbook('documents/Sektör Meslek Nace Kodları_2026.03.xlsx', data_only=True)
ws = wb['Sheet0']

profs_k = {}
for r in range(2, ws.max_row + 1):
    s_code = ws.cell(r, 1).value
    p_code = ws.cell(r, 3).value
    p_name = ws.cell(r, 4).value
    if s_code == 'K':
        if p_code not in profs_k:
            profs_k[p_code] = p_name

print("Sector K in Excel has", len(profs_k), "professions:")
for c, n in sorted(profs_k.items()):
    print(f"  {c}: {n}")
