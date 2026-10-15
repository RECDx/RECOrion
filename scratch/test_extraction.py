import openpyxl

wb = openpyxl.load_workbook('documents/Sektör Meslek Nace Kodları_2026.03.xlsx', data_only=True)
ws = wb['Sheet0']

sectors = {}
professions = {}

for r in range(2, ws.max_row + 1):
    c0 = ws.cell(r, 1).value
    c1 = ws.cell(r, 2).value
    c2 = ws.cell(r, 3).value
    c3 = ws.cell(r, 4).value
    
    sec_code = str(c0).strip() if c0 is not None else ""
    sec_name = str(c1).strip() if c1 is not None else ""
    prof_code = str(c2).strip() if c2 is not None else ""
    prof_name = str(c3).strip() if c3 is not None else ""
    
    if sec_code and sec_name:
        if sec_code not in sectors:
            sectors[sec_code] = sec_name
            
    if prof_code and prof_name and sec_code:
        if prof_code not in professions:
            professions[prof_code] = {
                "code": prof_code,
                "sector_code": sec_code,
                "name": prof_name
            }

print(f"Extracted {len(sectors)} Sectors:")
for sc, sn in sorted(sectors.items()):
    profs_in_sec = [p for p in professions.values() if p["sector_code"] == sc]
    print(f"  [{sc}] {sn} -> {len(profs_in_sec)} professions")
    if sc == 'K':
        for p in sorted(profs_in_sec, key=lambda x: x["code"]):
            print(f"      {p['code']}: {p['name']}")

print(f"\nTotal extracted professions: {len(professions)}")
